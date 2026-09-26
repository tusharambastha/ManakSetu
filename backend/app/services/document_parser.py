"""Document Parsing Service for Tender Documents (PDF and DOCX)
Extracts raw text and automatically isolates Technical Specifications,
Schedule of Requirements, or Scope of Work sections from full tender notices.
"""
import io
import re
from typing import Dict, Any, Optional
import pdfplumber
from docx import Document

class DocumentParserService:
    # Common headers demarcating technical specifications in Indian public tenders
    SPEC_HEADER_PATTERNS = [
        r"(?:section|schedule|annexure|part)\s*[-:]?\s*(?:v|iv|vi|iii|b|c|d)?\s*[-:]?\s*technical\s+specifications?",
        r"\btechnical\s+specifications?\b",
        r"technical\s+specifications?\s+(?:and|&)\s+requirements?",
        r"schedule\s+of\s+requirements?",
        r"bill\s+of\s+quantit(?:y|ies)\s*\(?boq\)?",
        r"scope\s+of\s+work\s*(?:and|&)?\s*technical\s+parameters?",
        r"technical\s+parameters?",
        r"item\s+specifications?",
        r"product\s+specifications?"
    ]

    # Delimiters marking the end of technical specifications
    END_SECTION_PATTERNS = [
        r"(?:section|schedule|annexure|part)\s*[-:]?\s*(?:vi|vii|viii|e|f)?\s*[-:]?\s*(?:commercial|financial|bidding|evaluation|terms|general\s+conditions)",
        r"financial\s+bid",
        r"commercial\s+terms\s+and\s+conditions",
        r"mode\s+of\s+payment",
        r"instructions\s+to\s+bidders",
        r"proforma\s+for",
        r"tender\s+form"
    ]

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> Dict[str, Any]:
        """Extract text from PDF and identify technical specification section"""
        full_text_pages = []
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for idx, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""
                    full_text_pages.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

        full_text = "\n".join(full_text_pages)
        isolated_spec, section_found = cls.extract_technical_section(full_text)

        return {
            "full_text": full_text,
            "extracted_spec": isolated_spec,
            "section_isolated": section_found,
            "total_pages": len(full_text_pages),
            "total_chars": len(full_text)
        }

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> Dict[str, Any]:
        """Extract text from DOCX and identify technical specification section"""
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text from tables (common in BOQ / tender schedules)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

        isolated_spec, section_found = cls.extract_technical_section(full_text)

        return {
            "full_text": full_text,
            "extracted_spec": isolated_spec,
            "section_isolated": section_found,
            "total_pages": 1,
            "total_chars": len(full_text)
        }

    @classmethod
    def extract_technical_section(cls, text: str) -> tuple[str, bool]:
        """Locates technical specifications section within full tender text"""
        if not text or len(text.strip()) < 30:
            return text.strip(), False

        # Attempt to find start header
        start_pos = -1
        best_match_header = ""
        for pattern in cls.SPEC_HEADER_PATTERNS:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                start_pos = match.start()
                best_match_header = match.group()
                break

        if start_pos == -1:
            # If no explicit header found, return first 4000 characters
            return text.strip()[:4000], False

        # If header found, search for closing delimiter after the header
        header_len = len(best_match_header)
        text_from_header = text[start_pos:]
        end_pos = len(text_from_header)

        for end_pat in cls.END_SECTION_PATTERNS:
            # Search after header_len to avoid matching within the header itself
            end_match = re.search(end_pat, text_from_header[header_len:], re.IGNORECASE)
            if end_match:
                end_pos = header_len + end_match.start()
                break

        extracted = text_from_header[:end_pos].strip()
        # Cap length if still massive (max 6000 chars)
        if len(extracted) > 6000:
            extracted = extracted[:6000]

        return extracted, True
