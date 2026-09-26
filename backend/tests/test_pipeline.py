"""Test Suite for Phase 3 & 4: NLP Extraction, Document Parsing & End-to-End Tender Analysis Pipeline"""
import io
import asyncio
from httpx import AsyncClient, ASGITransport
from docx import Document
from app.main import app
from app.services.nlp_extractor import NLPExtractorService
from app.services.document_parser import DocumentParserService

def run_async(coro):
    return asyncio.run(coro)

def test_multilingual_hindi_normalization():
    """Verify Hindi/Hinglish domain terms are normalized before retrieval"""
    hindi_query = "हमें साइट पर काम करने वाले मजदूरों के लिए सुरक्षा हेलमेट और सुरक्षा जूते चाहिए"
    normalized, was_detected = NLPExtractorService.normalize_multilingual_input(hindi_query)
    
    assert was_detected is True
    assert "safety helmet" in normalized.lower()
    assert "safety footwear" in normalized.lower() or "safety shoes" in normalized.lower()

def test_structured_parameter_extraction():
    """Verify technical parameters (voltage, strength, material) are cleanly extracted"""
    spec_text = "Supply of 1100V copper conductor PVC insulated cables with flame retardant properties for 50 Hz power system"
    reqs = NLPExtractorService.extract_structured_requirements(spec_text)
    
    assert reqs["inferred_sector"] == "Electrical & Power"
    assert "voltage" in reqs["technical_parameters"]
    assert any("1100" in v for v in reqs["technical_parameters"]["voltage"])
    assert "material" in reqs["technical_parameters"]
    assert any("copper" in m for m in reqs["technical_parameters"]["material"])

def test_document_section_extraction():
    """Verify isolation of technical specification from longer tender boilerplate"""
    tender_doc = """
    CENTRAL PUBLIC WORKS DEPARTMENT
    NOTICE INVITING TENDER (NIT No. 42/2026)
    Eligibility Criteria: Bidders must have annual turnover > 5 Crore.
    Earnest Money Deposit (EMD): Rs 50,000.
    
    SECTION IV: TECHNICAL SPECIFICATIONS
    Item 1: High Strength Deformed Steel Bars Fe 500D for seismic design.
    Item 2: Portland Pozzolana Cement PPC fly ash based meeting 33 MPa strength.
    Item 3: uPVC pipes class 4 for drinking water supply.
    
    SECTION V: COMMERCIAL TERMS AND CONDITIONS
    Payment terms: 90% against supply, 10% after commissioning.
    Arbitration jurisdiction: New Delhi.
    """
    spec, isolated = DocumentParserService.extract_technical_section(tender_doc)
    assert isolated is True
    assert "High Strength Deformed Steel Bars" in spec
    assert "COMMERCIAL TERMS" not in spec
    assert "Eligibility Criteria" not in spec

def test_end_to_end_text_analysis():
    """Verify /api/v1/analyze returns grounded results with anti-hallucination guarantees"""
    async def _test():
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            payload = {
                "text": "Technical requirements for Industrial Safety Helmets with 440V electrical insulation and chinstrap for shipyard workers",
                "top_k": 3,
                "session_id": "test-pipeline"
            }
            res = await client.post("/api/v1/analyze", json=payload)
            assert res.status_code == 200
            data = res.json()
            
            assert data["confident_match_found"] is True
            assert "extracted_requirements" in data
            assert data["extracted_requirements"]["inferred_sector"] == "PPE & Safety Equipment"
            
            top_rec = data["recommended_standards"][0]
            assert top_rec["standard_no"] == "IS 2925:1984"
            assert top_rec["status"] == "active"
            assert top_rec["certification"]["scheme"] == "BIS Product Certification"
            assert "IS 2925:1984" in top_rec["explanation"]
            assert len(top_rec["allied_standards"]) > 0
            assert "Zero synthetic standard numbers" in data["anti_hallucination_guarantee"]
    run_async(_test())

def test_end_to_end_upload_docx():
    """Verify /api/v1/analyze/upload handles uploaded DOCX file"""
    async def _test():
        # Create an in-memory DOCX file
        doc = Document()
        doc.add_heading("National Highway Authority of India - Tender", 0)
        doc.add_paragraph("General Instructions: E-tender mode only.")
        doc.add_heading("Technical Specifications", level=1)
        doc.add_paragraph("The contractor shall procure and supply TMT high strength deformed steel bars Fe 500D conforming to earthquake resistant requirements.")
        doc.add_heading("Commercial Conditions", level=1)
        doc.add_paragraph("Liquidated damages 0.5% per week of delay.")
        
        docx_bytes = io.BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            files = {"file": ("tender_spec.docx", docx_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            data_form = {"session_id": "test-pipeline"}
            res = await client.post("/api/v1/analyze/upload", files=files, data=data_form)
            assert res.status_code == 200
            data = res.json()
            
            assert data["confident_match_found"] is True
            assert data["document_metadata"]["filename"] == "tender_spec.docx"
            assert data["document_metadata"]["section_isolated"] is True
            
            # Should match IS 1786
            top_rec = data["recommended_standards"][0]
            assert top_rec["standard_no"] == "IS 1786:2008"
            assert top_rec["certification"]["scheme"] == "BIS Product Certification"
    run_async(_test())

def test_honest_no_match_case():
    """Verify that an unrelated query returns confident_match_found = False rather than forcing a hallucinated standard"""
    async def _test():
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            payload = {
                "text": "Fresh organic Alfonso mangoes from Ratnagiri farms export quality fruit boxes",
                "top_k": 3,
                "session_id": "test-pipeline"
            }
            res = await client.post("/api/v1/analyze", json=payload)
            assert res.status_code == 200
            data = res.json()
            assert data["confident_match_found"] is False
            assert "No confident Indian Standard match found" in data["pipeline_message"]
    run_async(_test())
